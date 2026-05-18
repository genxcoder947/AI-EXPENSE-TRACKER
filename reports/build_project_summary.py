from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("reports/ExpenseAI_Project_Summary.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(11, 37, 69)
MUTED = RGBColor(83, 100, 120)
INK = RGBColor(22, 33, 49)
LIGHT_BLUE_FILL = "E8EEF5"
LIGHT_GRAY_FILL = "F4F6F9"
PALE_GREEN_FILL = "ECFDF5"
PALE_YELLOW_FILL = "FFF7ED"
BORDER = "CBD5E1"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_in: float) -> None:
    width_dxa = int(width_in * 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_in)


def set_table_width(table, width_in: float = 6.5, indent_dxa: int = 120) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(width_in * 1440)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False


def set_table_borders(table, color: str = BORDER) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    return p


def add_callout(doc, label, body, fill=LIGHT_GRAY_FILL):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, 6.5, 120)
    set_table_borders(table, "D8DEE9")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_run_font(r, 10.5, NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(body)
    set_run_font(r2, 10.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, sum(widths), 120)
    set_table_borders(table)
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, 9.5, NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(value)
            set_run_font(r, 9.2, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "ExpenseAI Project Summary"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(run, 8.5, MUTED)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ExpenseAI Project Summary | Page ")
    set_run_font(run, 8.5, MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r = OxmlElement("w:r")
    r.append(fld_char1)
    r.append(instr)
    r.append(fld_char2)
    p._p.append(r)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_page_number_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in [
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def title_page(doc: Document) -> None:
    add_para(doc, "TECHNICAL PROJECT SUMMARY", size=10.5, color=BLUE, bold=True, after=4)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    title = p.add_run("ExpenseAI")
    set_run_font(title, 28, NAVY, bold=True)
    add_para(
        doc,
        "AI expense tracker: frontend, local storage, AI/OCR features, login flow, analytics, and macOS desktop app packaging.",
        size=13,
        color=MUTED,
        after=18,
    )

    meta = [
        ("Project type", "Local-first React expense tracking app"),
        ("Prepared on", "May 18, 2026"),
        ("Project folder", "/Users/aswinsanthosh/Downloads/src 3"),
        ("Desktop app", "/Users/aswinsanthosh/Desktop/ExpenseAI.app"),
        ("Primary build command", "npm run build"),
    ]
    add_table(doc, ["Field", "Details"], meta, [1.55, 4.95], header_fill=LIGHT_GRAY_FILL)

    add_callout(
        doc,
        "One-line summary",
        "ExpenseAI is currently a frontend-heavy, local-first app: the browser runs the React UI, local storage keeps app/user state, Tesseract.js reads receipt images, local rule-based AI suggests categories and insights, and a custom macOS app bundle launches the built site through a localhost server.",
        PALE_GREEN_FILL,
    )


def executive_summary(doc: Document) -> None:
    add_heading(doc, "1. Executive Summary", 1)
    add_bullet(doc, "ExpenseAI is a smart expense tracker built as a React single-page application.")
    add_bullet(doc, "The app supports login/sign-up, manual expense entry, optional notes, receipt scanning, AI category suggestions, analytics cards, spending charts, edit/delete controls, and dark/light theme.")
    add_bullet(doc, "There is no separate cloud backend or production database in the current project. The app uses browser-side state and localStorage for the local demo experience.")
    add_bullet(doc, "Receipt scanning uses Tesseract.js OCR in the browser and a custom parser that detects totals, merchants, and item-level bill details.")
    add_bullet(doc, "The macOS desktop app is a custom .app bundle that starts a localhost server and opens the built site at http://127.0.0.1:8000/index.html.")


def tech_stack(doc: Document) -> None:
    add_heading(doc, "2. Technology Stack", 1)
    rows = [
        ("Frontend framework", "React 18.3.1 with TypeScript", "src/App.tsx, src/main.tsx, src/components/*", "Builds the interactive dashboard and expense workflow."),
        ("Build tool", "Vite 6 with @vitejs/plugin-react", "vite.config.ts, package.json", "Runs the dev server and creates static production assets in dist/."),
        ("Styling", "Tailwind CSS 3, PostCSS, custom CSS", "tailwind.config.js, postcss.config.js, src/index.css", "Responsive UI, theme switching, chart animations, dark-mode polish."),
        ("Icons", "lucide-react", "Buttons, cards, scan panel, chart controls", "Adds consistent UI icons such as Sparkles, Camera, Trash, Pencil, Calendar."),
        ("AI logic", "Local rule-based categorization and insight engine", "src/utils/aiEngine.ts", "Suggests categories, confidence, reasons, practical actions, and spending insights."),
        ("OCR", "Tesseract.js 6", "src/components/ReceiptScanPanel.tsx", "Reads text from receipt photos/screenshots or live camera captures."),
        ("Storage", "Browser localStorage plus React state", "src/App.tsx, src/lib/auth.ts", "Stores demo users/current user/theme; expenses are initialized from localStorage and managed in app state."),
        ("Desktop app", "macOS .app bundle + LaunchAgent + Python/Node HTTP server", "ExpenseAI.app/Contents/*", "Lets the project open like a normal Mac app without manually running terminal commands."),
    ]
    add_table(doc, ["Area", "What is used", "Where", "Purpose"], rows, [1.35, 1.65, 1.65, 1.85])


def architecture(doc: Document) -> None:
    add_heading(doc, "3. Runtime Architecture", 1)
    add_callout(
        doc,
        "Backend status",
        "Current project backend: no traditional backend API. The app is a static React build. In development, Vite serves it. In the desktop app, a local HTTP server serves the built files from ExpenseAI.app/Contents/Resources/site.",
        PALE_YELLOW_FILL,
    )
    add_number(doc, "User opens ExpenseAI.app or the Vite dev URL.")
    add_number(doc, "The browser loads index.html and bundled React assets.")
    add_number(doc, "React mounts App.tsx, checks localStorage for currentUser, theme, and existing expense data.")
    add_number(doc, "If no user is active, Login.tsx shows the login/sign-up form.")
    add_number(doc, "After login, the dashboard renders stats, expense form, expense list, insights, and charts.")
    add_number(doc, "Receipt images are processed in the browser by Tesseract.js; parsed results are copied into the expense form.")
    add_number(doc, "The app computes AI suggestions, insights, category totals, and chart slices locally from the expense array.")


def frontend_features(doc: Document) -> None:
    add_heading(doc, "4. Frontend and UI Features", 1)
    rows = [
        ("Main dashboard", "App.tsx", "Controls app state, login gate, theme toggle, header, layout grid, stats, form, list, insights, and charts."),
        ("Manual add expense", "ExpenseForm.tsx", "Inputs for description, amount, date, category, and optional notes. Creates an Expense object with generated id."),
        ("Optional notes", "ExpenseForm.tsx, ExpenseList.tsx", "Notes are optional and capped at 700 characters. Receipt scans can auto-fill notes with item details."),
        ("Edit/delete expense", "ExpenseList.tsx", "Each expense row supports inline edit for description, amount, date, category, notes, plus delete."),
        ("Stats details", "StatsOverview.tsx", "Clickable stat cards open detail modals for total spent, this month, and AI categorized expenses."),
        ("Charts", "SpendingChart.tsx, CategoryPieChart.tsx", "Shows category totals, pie chart, time ranges, search filters, and animated full chart modal."),
        ("Themes", "App.tsx, index.css", "Class-based light/dark theme; recent dark-mode CSS makes cards, inputs, hovers, and panels consistent."),
        ("Reusable UI", "components/ui/*", "Small Button, Card, Input, Label primitives shared across screens."),
    ]
    add_table(doc, ["Feature", "File", "What it does"], rows, [1.35, 1.65, 3.5])


def data_and_auth(doc: Document) -> None:
    add_heading(doc, "5. Login, Storage, and Data Model", 1)
    add_heading(doc, "Login Page", 2)
    add_bullet(doc, "Login.tsx handles both sign-in and account creation from the same card UI.")
    add_bullet(doc, "It validates that the email contains @ and the password has at least 6 characters.")
    add_bullet(doc, "The form uses the shared Input component, now with dark text/caret fixes so typed text remains visible.")

    add_heading(doc, "Local Demo Auth", 2)
    rows = [
        ("users", "localStorage key", "Stores demo user records by email/username."),
        ("currentUser", "localStorage key", "Stores the active signed-in user."),
        ("theme", "localStorage key", "Stores light/dark theme preference."),
        ("expenses", "localStorage key read in App.tsx", "The app initializes expenses from this key, while current expense updates are handled through React state."),
    ]
    add_table(doc, ["Data", "Storage", "Purpose"], rows, [1.4, 1.8, 3.3], header_fill=LIGHT_GRAY_FILL)
    add_callout(
        doc,
        "Important security note",
        "This login system is for a local/demo project. Passwords are stored directly in localStorage, so it is not safe for real users. A production version should use a backend auth service, hashed passwords, secure sessions, and a real database.",
        PALE_YELLOW_FILL,
    )

    add_heading(doc, "Expense Data Shape", 2)
    add_bullet(doc, "Expense fields: id, description, amount, category, date, optional notes, and aiSuggested.")
    add_bullet(doc, "Dates are stored as YYYY-MM-DD strings and displayed with Intl.DateTimeFormat.")
    add_bullet(doc, "Currency formatting uses Intl.NumberFormat with en-IN and INR.")


def ai_ocr_features(doc: Document) -> None:
    add_heading(doc, "6. AI and Receipt Scanning Features", 1)
    add_heading(doc, "AI Category Suggestions", 2)
    add_bullet(doc, "suggestCategory(description) tokenizes the description and compares it with category keyword lists.")
    add_bullet(doc, "Supported categories include Food & Dining, Transportation, Shopping, Entertainment, Bills & Utilities, Healthcare, Education, Travel, and Other.")
    add_bullet(doc, "The suggestion includes a category, confidence score, reason, matched signals, and a practical action tip.")
    add_bullet(doc, "Example: Zomato or Swiggy signals map to Food & Dining with an action about separating eating-out spend from groceries.")

    add_heading(doc, "AI Insights", 2)
    add_bullet(doc, "generateInsights calculates category concentration, weekly daily average, projected monthly spend, and practical spending tips.")
    add_bullet(doc, "calculateCategoryTotals computes totals and percentages for charts.")

    add_heading(doc, "Receipt OCR and Item Details", 2)
    rows = [
        ("Image input", "Choose photo, device camera/gallery, or live webcam capture."),
        ("OCR engine", "Tesseract.js creates an English OCR worker and reads text from the selected image."),
        ("Parser", "receiptParser.ts detects merchant names, total/payment lines, rupee/dollar amounts, and line items."),
        ("Auto-filled fields", "Description, total amount, optional notes, item count, confidence, and warnings."),
        ("Notes output", "Scanned item names, quantities when detected, individual item amounts, and receipt total."),
        ("Verification", "The UI warns the user to verify OCR results before saving because real receipts can be noisy."),
    ]
    add_table(doc, ["Part", "How it works"], rows, [1.5, 5.0])
    add_callout(
        doc,
        "AI type used",
        "The current project does not call a cloud LLM. The 'AI' behavior is local deterministic logic: keyword scoring, confidence calculation, practical advice rules, and OCR-assisted receipt parsing.",
        PALE_GREEN_FILL,
    )


def analytics_and_charts(doc: Document) -> None:
    add_heading(doc, "7. Analytics, Charts, and Animations", 1)
    add_bullet(doc, "StatsOverview shows Total Spent, This Month, and AI Categorized. Each card opens a modal with matching expense details.")
    add_bullet(doc, "SpendingChart shows the top category totals as progress bars.")
    add_bullet(doc, "CategoryPieChart renders a custom SVG donut/pie chart with stable category colors.")
    add_bullet(doc, "The full chart modal supports Day, Last 24h, Month, and Year filters plus category/description search.")
    add_bullet(doc, "Chart animations in index.css include modal fade-in, panel scale-in, pie pop-in, slice drawing, and center-text fade.")
    add_bullet(doc, "Reduced-motion support disables chart animations for users who prefer less motion.")


def desktop_app(doc: Document) -> None:
    add_heading(doc, "8. macOS Desktop App Packaging", 1)
    add_callout(
        doc,
        "How the app became a Mac app",
        "The project was not converted into Electron. Instead, the Vite build was copied into a normal macOS .app bundle. The bundle starts a tiny local server and opens the app URL in the browser, so the user can launch ExpenseAI by double-clicking the desktop app.",
        PALE_GREEN_FILL,
    )
    rows = [
        ("ExpenseAI.app/Contents/Info.plist", "Defines bundle name, executable, identifier, version, package type, and LSUIElement."),
        ("ExpenseAI.app/Contents/MacOS/ExpenseAI", "zsh launcher that checks localhost, writes/starts a LaunchAgent, waits for the server, and opens the app URL."),
        ("ExpenseAI.app/Contents/Resources/site", "Contains the built Vite site: index.html plus hashed JS/CSS assets."),
        ("ExpenseAI.app/Contents/Resources/server.mjs", "Node fallback static server if Python 3 is unavailable."),
        ("~/Library/LaunchAgents/com.local.expenseai.server.plist", "Generated by the launcher to run Python http.server or the Node fallback."),
        ("http://127.0.0.1:8000/index.html", "Stable local URL used by the desktop app. Keeping the same origin helps preserve browser localStorage data."),
    ]
    add_table(doc, ["Piece", "Purpose"], rows, [2.55, 3.95])

    add_heading(doc, "Desktop Launch Flow", 2)
    add_number(doc, "User double-clicks ExpenseAI.app on the Desktop.")
    add_number(doc, "The shell launcher checks whether http://127.0.0.1:8000/index.html is reachable.")
    add_number(doc, "If the server is not running, the launcher writes a LaunchAgent plist and starts it.")
    add_number(doc, "The LaunchAgent serves the built site from Contents/Resources/site.")
    add_number(doc, "The launcher opens the local URL in the default browser.")


def build_and_files(doc: Document) -> None:
    add_heading(doc, "9. Build, Run, and Important Files", 1)
    rows = [
        ("npm run dev", "Starts the Vite development server for local development."),
        ("npm run build", "Runs TypeScript checks and creates the production build in dist/."),
        ("npm run preview", "Serves the built Vite output for preview."),
        ("cp -R dist/. ExpenseAI.app/Contents/Resources/site/", "Syncs the production build into the local app bundle."),
        ("cp -R dist/. /Users/aswinsanthosh/Desktop/ExpenseAI.app/Contents/Resources/site/", "Syncs the build into the Desktop app copy."),
    ]
    add_table(doc, ["Command", "Meaning"], rows, [2.7, 3.8], header_fill=LIGHT_GRAY_FILL)

    add_heading(doc, "Key Source Files", 2)
    rows = [
        ("src/App.tsx", "Application shell, state, login gate, theme toggle, layout."),
        ("src/components/ExpenseForm.tsx", "Manual expense entry, optional notes, AI suggestion UI, receipt scan handoff."),
        ("src/components/ReceiptScanPanel.tsx", "Photo upload, live camera, webcam capture, Tesseract.js OCR."),
        ("src/utils/receiptParser.ts", "Receipt amount, merchant, item, confidence, and warning extraction."),
        ("src/utils/aiEngine.ts", "Category suggestions, actionable tips, insights, category totals."),
        ("src/components/ExpenseList.tsx", "Expense display, optional notes display, edit and delete controls."),
        ("src/components/CategoryPieChart.tsx", "Time-filtered category pie chart, modal, search, SVG rendering."),
        ("src/lib/auth.ts", "Local demo sign-up/login/current-user helpers."),
        ("src/index.css", "Tailwind entry, chart animations, dark-theme overrides."),
    ]
    add_table(doc, ["File", "Role"], rows, [2.1, 4.4])


def limitations(doc: Document) -> None:
    add_heading(doc, "10. Current Limitations and Production Improvements", 1)
    add_bullet(doc, "Use a real backend and database for production, such as Node/Express or serverless APIs plus Postgres/Firebase/Supabase.")
    add_bullet(doc, "Replace localStorage demo auth with secure authentication, hashed passwords, sessions/tokens, and protected user-specific data.")
    add_bullet(doc, "Persist expense updates reliably per user. Current source initializes expenses from localStorage but manages changes in React state.")
    add_bullet(doc, "Improve OCR with receipt image preprocessing, language/worker asset bundling, confidence scoring per field, and manual correction history.")
    add_bullet(doc, "If true generative AI is needed, add an API-backed assistant for richer natural-language suggestions, budget plans, and anomaly detection.")
    add_bullet(doc, "Consider packaging as Electron/Tauri later if a fully native app window, menu bar, offline asset control, and stronger desktop integration are required.")


def main() -> None:
    doc = Document()
    configure_document(doc)
    title_page(doc)
    executive_summary(doc)
    tech_stack(doc)
    architecture(doc)
    frontend_features(doc)
    data_and_auth(doc)
    ai_ocr_features(doc)
    analytics_and_charts(doc)
    desktop_app(doc)
    build_and_files(doc)
    limitations(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
